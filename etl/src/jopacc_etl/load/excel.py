"""Write one workbook per report.

Each workbook is self-contained: the figures, the things that did not parse,
a page-by-page record of how the file was read, the report's own metadata, and
the validation findings. A reader who opens one of these should be able to
answer "where did this number come from?" without going back to the pipeline.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List

from .curated import curate_measurements

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="13293D")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
SEVERITY_FILL = {
    "error": PatternFill("solid", fgColor="F8D7DA"),
    "warning": PatternFill("solid", fgColor="FFF3CD"),
    "info": PatternFill("solid", fgColor="E2E8F0"),
}

MEASUREMENT_COLUMNS = [
    ("record_id", 26), ("source_pdf", 34), ("report_type", 12),
    ("reporting_period", 16), ("source_page", 11), ("section_key", 17),
    ("section_title", 30), ("panel_index", 11), ("panel_title", 38),
    ("metric", 32), ("metric_source", 14),
    ("label", 24), ("label_period", 13), ("raw_value", 14),
    ("value_numeric", 18), ("unit", 9), ("scale_suffix", 12),
    ("extraction_method", 17), ("ocr_layer", 11), ("confidence", 11),
    ("ocr_repaired", 13), ("ocr_alternatives", 20), ("x_pt", 9), ("y_pt", 9),
    ("note", 34), ("extracted_at_utc", 20),
]


def _write_sheet(book: Workbook, title: str, columns: List[tuple],
                 rows: List[dict], table_name: str) -> None:
    sheet = book.create_sheet(title)
    headers = [c[0] for c in columns]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        sheet.append([row.get(h) for h in headers])

    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1", showRowStripes=True)
        sheet.add_table(table)


def _metadata_rows(result) -> List[dict]:
    rec = result.record
    counts = result.counts
    pairs = [
        ("Source PDF", rec.filename),
        ("Report type", rec.report_type),
        ("Report type determined from", rec.report_type_source),
        ("Reporting period", rec.period_key),
        ("Reporting period determined from", rec.period_source),
        ("Period stated in filename", rec.filename_period or "not stated"),
        ("Within configured scope window", rec.in_scope),
        ("Pages", rec.page_count),
        ("Pages with a text layer", rec.text_pages),
        ("Pages requiring OCR", rec.outlined_pages),
        ("Canvas size (pt)", f"{rec.canvas_w_pt} x {rec.canvas_h_pt}"),
        ("PDF created", rec.pdf_created or "not recorded"),
        ("PDF last modified", rec.pdf_modified or "not recorded"),
        ("PDF producer", rec.producer or "not recorded"),
        ("File size (bytes)", rec.size_bytes),
        ("SHA-256 of file", rec.sha256),
        ("Processing status", result.status),
        ("Pages read successfully", counts["pages_ok"]),
        ("Pages failed", counts["pages_failed"]),
        ("Pages read natively", counts["pages_native"]),
        ("Pages read by OCR", counts["pages_ocr"]),
        ("Measurements extracted", counts["measurements"]),
        ("Tokens with digits left unparsed", counts["unparsed_tokens"]),
    ]
    return [{"property": k, "value": v} for k, v in pairs]


def _section_rows(result) -> List[dict]:
    return [{
        "section_key": s.key, "section_title": s.title,
        "start_page": s.start_page, "end_page": s.end_page,
        "pages": s.end_page - s.start_page + 1,
        "recognised": s.recognised,
    } for s in result.sections]


def write_workbook(result, cfg: dict, output_root: Path) -> Path:
    book = Workbook()
    book.remove(book.active)
    sheets = cfg["load"]["sheets"]

    curated_rows = curate_measurements(result.measurements, result.record)
    _write_sheet(book, "Curated",
                 [("report_type", 12), ("reporting_period", 16), ("system", 16),
                  ("system_title", 34), ("metric", 34), ("dimension", 22),
                  ("dimension_period", 14), ("period", 14), ("value", 16),
                  ("unit", 9), ("raw_value", 14), ("confidence", 11),
                  ("extraction_method", 16), ("flagged", 9),
                  ("source_pdf", 32), ("source_page", 12)],
                 curated_rows, "Curated")
    _write_sheet(book, sheets["measurements"], MEASUREMENT_COLUMNS,
                 result.measurements, "Measurements")
    _write_sheet(book, "Sections",
                 [("section_key", 18), ("section_title", 46), ("start_page", 12),
                  ("end_page", 11), ("pages", 9), ("recognised", 12)],
                 _section_rows(result), "Sections")
    _write_sheet(book, sheets["unparsed"],
                 [("source_pdf", 34), ("source_page", 11), ("section_key", 17),
                  ("raw_text", 20), ("extraction_method", 17), ("confidence", 11),
                  ("x_pt", 9), ("y_pt", 9), ("reason", 52)],
                 result.unparsed, "Unparsed")
    _write_sheet(book, sheets["pages"],
                 [("page", 8), ("status", 10), ("method", 10), ("token_count", 13),
                  ("ornament_dropped", 18), ("rotated_tokens", 16),
                  ("width_pt", 11), ("height_pt", 11), ("error", 50)],
                 result.page_log, "PageLog")
    _write_sheet(book, sheets["metadata"],
                 [("property", 38), ("value", 66)], _metadata_rows(result), "Metadata")

    findings = [f.as_dict() for f in result.findings]
    _write_sheet(book, sheets["validation"],
                 [("severity", 11), ("check", 28), ("page", 8),
                  ("record_id", 26), ("detail", 96)],
                 findings, "Validation")

    sheet = book[sheets["validation"]]
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        fill = SEVERITY_FILL.get(row[0].value)
        if fill:
            for cell in row:
                cell.fill = fill

    rec = result.record
    name = cfg["load"]["filename_template"].format(
        report_type=rec.report_type or "unclassified",
        period_key=(rec.period_key or Path(rec.filename).stem).replace("-", "_"),
    )
    destination = output_root
    if cfg["load"]["subdirectory_by_type"]:
        destination = destination / (rec.report_type or "unclassified")
    destination.mkdir(parents=True, exist_ok=True)
    path = destination / name
    book.save(path)
    log.info("wrote %s (%d measurements)", path.name, len(result.measurements))
    return path
