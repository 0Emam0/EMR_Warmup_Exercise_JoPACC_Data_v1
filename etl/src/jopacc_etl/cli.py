"""Command line entry point.

    python -m jopacc_etl inventory      # register PDFs, report scope and gaps
    python -m jopacc_etl run            # extract, validate and write workbooks
    python -m jopacc_etl run --only monthly --period 2026-08
    python -m jopacc_etl run --failed   # retry only what failed last time
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

from .config import load_config
from .discovery import inventory as inv
from .load.excel import write_workbook
from .load.curated import write_master_dataset
from .logging_setup import setup
from .pipeline import safe_process
from .validation.checks import summarise

log = logging.getLogger("jopacc_etl")


def _state_paths(cfg):
    state = Path(cfg["paths"]["state_dir"])
    state.mkdir(parents=True, exist_ok=True)
    return state / "inventory.json", state / "run_state.json"


def cmd_inventory(cfg, args) -> int:
    inventory_path, _ = _state_paths(cfg)
    records = inv.discover(cfg)
    inv.save(records, inventory_path)

    dupes = inv.detect_duplicates(records)
    gaps = inv.missing_periods(records, cfg)

    print(f"\nRegistered {len(records)} PDF(s) from {cfg['paths']['input_dir']}\n")
    by_type = {}
    for rec in records:
        by_type.setdefault(rec.report_type or "unclassified", []).append(rec)
    for report_type, group in sorted(by_type.items()):
        in_scope = sum(1 for r in group if r.in_scope)
        ocr = sum(1 for r in group if r.outlined_pages > r.text_pages)
        print(f"  {report_type:12s} {len(group):3d} file(s), {in_scope} in scope, "
              f"{ocr} needing OCR")

    if dupes["identical_files"]:
        print("\n  Identical files:")
        for names in dupes["identical_files"].values():
            print("    ", ", ".join(names))
    if dupes["period_collisions"]:
        print("\n  More than one file for the same period:")
        for period, names in dupes["period_collisions"].items():
            print(f"     {period}: {', '.join(names)}")

    for report_type, missing in gaps.items():
        if missing:
            print(f"\n  Missing {report_type} period(s): {', '.join(missing)}")

    problems = [r for r in records if r.issues]
    if problems:
        print("\n  Files with issues:")
        for rec in problems:
            print(f"     {rec.filename}: {'; '.join(rec.issues)}")
    print()
    return 0


def _select(records, args):
    chosen = records
    if args.only:
        chosen = [r for r in chosen if r.report_type == args.only]
    if args.period:
        chosen = [r for r in chosen if r.period_key == args.period]
    if not getattr(args, "include_out_of_scope", False):
        chosen = [r for r in chosen if r.in_scope is not False]
    return chosen


def cmd_run(cfg, args) -> int:
    inventory_path, run_state_path = _state_paths(cfg)
    records = inv.load(inventory_path)
    if not records:
        log.info("no inventory found; discovering first")
        records = inv.discover(cfg)
        inv.save(records, inventory_path)

    previous = {}
    if run_state_path.exists():
        previous = json.loads(run_state_path.read_text())

    targets = _select(records, args)

    if args.failed:
        targets = [r for r in targets
                   if previous.get(r.filename, {}).get("status") not in ("ok",)]
    elif not args.force:
        # Skip files already processed successfully, unless their content hash
        # has changed. JoPACC replaces published files without renaming them.
        kept = []
        for rec in targets:
            seen = previous.get(rec.filename)
            if not seen or seen.get("status") != "ok":
                kept.append(rec)
            elif cfg["runtime"]["detect_replaced_files"] and seen.get("sha256") != rec.sha256:
                log.warning("%s has been replaced since the last run; reprocessing",
                            rec.filename)
                kept.append(rec)
        skipped = len(targets) - len(kept)
        if skipped:
            log.info("skipping %d unchanged file(s) already processed", skipped)
        targets = kept

    if not targets:
        print("Nothing to process.")
        return 0

    log.info("processing %d report(s) with %d worker(s)",
             len(targets), cfg["runtime"]["max_workers"])
    output_root = Path(cfg["paths"]["output_dir"])
    results = []

    state = dict(previous)
    summary = []

    def finish(result):
        """Write the workbook and commit state for one report.

        State is written after every report rather than at the end of the run.
        A long OCR run that is interrupted then resumes where it stopped
        instead of starting over.
        """
        try:
            path = write_workbook(result, cfg, output_root)
        except Exception as exc:
            log.error("could not write workbook for %s: %s", result.record.filename, exc)
            result.status = "failed"
            result.error = f"workbook write failed: {exc}"
            path = None
        counts = result.counts
        state[result.record.filename] = {
            "status": result.status,
            "sha256": result.record.sha256,
            "workbook": str(path) if path else "",
            **counts,
        }
        run_state_path.write_text(json.dumps(state, indent=1))
        summary.append((result, counts))
        log.info("[%d/%d] %s -> %s", len(summary), len(targets),
                 result.record.filename, result.status)

    workers = max(1, int(args.workers or cfg["runtime"]["max_workers"]))
    if workers == 1:
        for rec in targets:
            finish(safe_process(rec, cfg))
    else:
        with ProcessPoolExecutor(max_workers=workers) as pool:
            futures = [pool.submit(safe_process, rec, cfg) for rec in targets]
            for future in as_completed(futures):
                finish(future.result())
    _print_quality_report(summary, cfg)
    return 0 if all(r.status != "failed" for r, _ in summary) else 1


def _print_quality_report(summary, cfg) -> None:
    totals = {"pages_total": 0, "pages_ok": 0, "pages_failed": 0,
              "pages_native": 0, "pages_ocr": 0, "measurements": 0,
              "unparsed_tokens": 0}
    findings = {"error": 0, "warning": 0, "info": 0}
    ok = partial = failed = 0

    for result, counts in summary:
        for key in totals:
            totals[key] += counts[key]
        for level, count in summarise(result.findings).items():
            findings[level] = findings.get(level, 0) + count
        ok += result.status == "ok"
        partial += result.status == "partial"
        failed += result.status == "failed"

    print("\n" + "=" * 62)
    print("DATA QUALITY REPORT")
    print("=" * 62)
    print(f"  Reports processed        {len(summary)}")
    print(f"    fully successful       {ok}")
    print(f"    partial                {partial}")
    print(f"    failed                 {failed}")
    print(f"  Pages seen               {totals['pages_total']}")
    print(f"    read successfully      {totals['pages_ok']}")
    print(f"    failed                 {totals['pages_failed']}")
    print(f"    read natively          {totals['pages_native']}")
    print(f"    read by OCR            {totals['pages_ocr']}")
    print(f"  Measurements loaded      {totals['measurements']}")
    print(f"  Digit tokens unparsed    {totals['unparsed_tokens']}")
    print(f"  Validation findings      {findings['error']} error, "
          f"{findings['warning']} warning, {findings['info']} info")
    print("=" * 62)
    for result, counts in summary:
        if result.status != "ok":
            print(f"  ! {result.record.filename}: {result.status} -- {result.error}")
    print()


def cmd_dataset(cfg, args) -> int:
    """Combine the Curated sheet of every workbook on disk into one file."""
    import openpyxl

    output_root = Path(cfg["paths"]["output_dir"])
    rows = []
    for path in sorted(output_root.rglob("*.xlsx")):
        if path.name == "jopacc_dataset.xlsx":
            continue
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Curated" not in book.sheetnames:
            continue
        sheet = book["Curated"]
        headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(headers, row)))
        book.close()

    if not rows:
        print("No workbooks with a Curated sheet found under", output_root)
        return 1

    destination = output_root / "jopacc_dataset.xlsx"
    write_master_dataset(rows, destination)
    print(f"Combined {len(rows)} rows from "
          f"{len({(r['report_type'], r['reporting_period']) for r in rows})} "
          f"report(s) into {destination}")
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(prog="jopacc_etl",
                                     description="ETL for JoPACC payment system reports")
    parser.add_argument("--config", default="config/config.yaml")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("inventory", help="register PDFs and report scope coverage")
    sub.add_parser("dataset", help="combine every workbook's Curated sheet into one file")

    run = sub.add_parser("run", help="extract, validate and write workbooks")
    run.add_argument("--only", choices=["monthly", "quarterly", "yearly"])
    run.add_argument("--period", help="a single period key, e.g. 2026-08 or 2026-Q1")
    run.add_argument("--failed", action="store_true",
                     help="reprocess only reports that did not succeed last run")
    run.add_argument("--force", action="store_true",
                     help="reprocess everything, ignoring previous run state")
    run.add_argument("--include-out-of-scope", action="store_true")
    run.add_argument("--workers", type=int)

    args = parser.parse_args(argv)
    cfg = load_config(args.config)
    log_file = setup(cfg)
    log.info("configuration loaded from %s; logging to %s", args.config, log_file)

    if args.command == "inventory":
        return cmd_inventory(cfg, args)
    if args.command == "dataset":
        return cmd_dataset(cfg, args)
    return cmd_run(cfg, args)


if __name__ == "__main__":
    sys.exit(main())
